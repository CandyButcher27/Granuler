import io
import os
import re
import secrets
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from fastapi import Depends, FastAPI, HTTPException, Query, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, Response
from fastapi.security import HTTPBasic, HTTPBasicCredentials
from pydantic import BaseModel, Field

from .llm import (
    generate_pillar_content,
    generate_executive,
    generate_context,
    generate_findings,
    generate_risks,
    generate_deep_dives,
    generate_architecture,
    generate_plan,
    generate_granuler,
    generate_prior_work,
    generate_quick_wins,
    generate_risk_register,
    generate_proposal,
    extract_from_notes,
)
from .image_brief import image_brief_pdf, slots_from_deck
from .pdf_generator import RENDERERS as PDF_RENDERERS
from .deck import (
    PILLAR_COUNT,
    SUBTOPICS_PER_PILLAR,
    band_for,
    overall_score,
    pillar_score,
)
from .slides import build_deck

# Single-user gate. One username and password, supplied by the environment.
# `/health` stays open so the host's health check does not need credentials.
_AUTH_USER = os.environ.get("GRANULER_USER", "")
_AUTH_PASSWORD = os.environ.get("GRANULER_PASSWORD", "")
_OPEN_PATHS = {"/health"}
_basic = HTTPBasic(auto_error=False)


def _require_login(request: Request, creds: HTTPBasicCredentials = Depends(_basic)) -> None:
    """Reject anyone who is not the single configured user.

    Fails closed: with no password configured the service refuses everything
    rather than serving the assessment tool to the open internet. Both halves
    are always compared so a wrong username costs the same time as a wrong
    password.
    """
    if request.url.path in _OPEN_PATHS:
        return
    if not _AUTH_PASSWORD:
        raise HTTPException(
            status_code=503,
            detail="Server is not configured for sign-in. Set GRANULER_USER and GRANULER_PASSWORD.",
        )
    user_ok = secrets.compare_digest((creds.username if creds else ""), _AUTH_USER)
    password_ok = secrets.compare_digest((creds.password if creds else ""), _AUTH_PASSWORD)
    if not (user_ok and password_ok):
        raise HTTPException(
            status_code=401,
            detail="Incorrect username or password",
            headers={"WWW-Authenticate": 'Basic realm="Granuler"'},
        )


app = FastAPI(title="Granuler Report API", dependencies=[Depends(_require_login)])

# Client identity is never sent to the LLM. Prompts use this placeholder;
# _restore_name() swaps it back to the real company name in the LLM's JSON
# response before it reaches the deck or the API caller.
_LLM_CLIENT_LABEL = "the client company"
# Case-insensitive: the LLM capitalises the placeholder at sentence start.
_LLM_CLIENT_RE = re.compile(re.escape(_LLM_CLIENT_LABEL), re.IGNORECASE)


def _restore_name(obj, company_name: str):
    if isinstance(obj, str):
        return _LLM_CLIENT_RE.sub(lambda _: company_name, obj)
    if isinstance(obj, list):
        return [_restore_name(v, company_name) for v in obj]
    if isinstance(obj, dict):
        return {k: _restore_name(v, company_name) for k, v in obj.items()}
    return obj

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


class SubtopicIn(BaseModel):
    subtopic: str
    score: int = Field(ge=1, le=5)
    weighted_marks: float = 0.0
    impact: str = "Medium"
    priority: str = "Medium"
    current_state_notes: str = ""
    evidence: str = ""
    recommended_action: str = ""
    owner: str = ""
    timeline: str = ""


class PillarIn(BaseModel):
    pillar: str
    subtopics: list[SubtopicIn]


class ReportRequest(BaseModel):
    company_name: str
    industry: str = ""
    assessment_date: str = ""
    assessor: str = "Ravi Kajaria"
    business_goals: str = ""
    pain_points: str = ""
    revenue_range: str = ""
    employee_count: str = ""
    locations: str = ""
    core_systems: str = ""
    major_risks: str = ""
    key_stakeholders: str = ""
    priority_areas: str = ""
    budget_appetite: str = ""
    change_readiness: str = ""
    founder_dependency: str = ""
    products: str = ""
    industries_served: str = ""
    granuler_location: str = "Mumbai"
    savings_identified: str = ""
    # Work Granuler has already delivered for this client. Empty for a new
    # client, which drops the "progress delivered" slide from the deck rather
    # than have the LLM invent a track record.
    prior_work: str = ""
    pillars: list[PillarIn]


DEMO_HTML_PATH = Path(__file__).resolve().parent.parent / "demo.html"
LOGO_PATH = Path(__file__).resolve().parent.parent / "assets" / "granuler_logo.png"


@app.get("/")
def demo():
    return FileResponse(DEMO_HTML_PATH)


@app.get("/health")
def health():
    return {"status": "ok"}


@app.get("/logo.png")
def logo():
    return FileResponse(LOGO_PATH, media_type="image/png")


@app.post("/image-brief")
async def image_brief(request: Request):
    """The prompt sheet for a generated deck's numbered image placeholders.

    Takes the deck back rather than the form data: the API keeps no state, and
    a deck's own speaker notes are the only record of which slots it opened
    that cannot drift from the file the assessor is actually holding.
    """
    from pptx import Presentation

    body = await request.body()
    if not body:
        raise HTTPException(status_code=422, detail="Post the generated .pptx as the body")
    try:
        slots = slots_from_deck(Presentation(io.BytesIO(body)))
    except Exception:
        raise HTTPException(status_code=422, detail="Body is not a readable .pptx")
    company = request.query_params.get("company_name", "")
    name = "Granuler_Image_Brief" if not company else \
        f"{re.sub(r'[^A-Za-z0-9]+', '_', company).strip('_')}_Image_Brief"
    return Response(
        content=image_brief_pdf(slots, company=company),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{name}.pdf"'},
    )


@app.post("/generate-report")
def generate(req: ReportRequest):
    if len(req.pillars) != PILLAR_COUNT:
        raise HTTPException(status_code=422, detail=f"Exactly {PILLAR_COUNT} pillars required")

    pillars_raw = [p.model_dump() for p in req.pillars]
    intake = req.model_dump(exclude={"pillars"})

    score = overall_score(pillars_raw, SUBTOPICS_PER_PILLAR)
    band = band_for(score)[0]
    pillar_summaries = [
        {"name": p["pillar"], "score": pillar_score(p["subtopics"], SUBTOPICS_PER_PILLAR)}
        for p in pillars_raw
    ]

    # Every narrative prompt needs the same situation brief. The company's real
    # name is never part of it - see _LLM_CLIENT_LABEL.
    brief = dict(
        company_name=_LLM_CLIENT_LABEL,
        industry=req.industry,
        business_goals=req.business_goals,
        pain_points=req.pain_points,
        core_systems=req.core_systems,
        major_risks=req.major_risks,
        overall_score=score,
        maturity_band=band,
        pillar_summaries=pillar_summaries,
    )

    # Eight or nine section calls plus one per pillar. Run sequentially that is 60-90s;
    # fanned out over threads the wall clock is roughly the slowest single call.
    # litellm.completion is blocking HTTP, so threads are the right primitive.
    jobs: dict[str, tuple] = {
        "executive": (generate_executive, brief),
        "context": (generate_context, dict(
            brief,
            locations=req.locations,
            products=req.products,
            industries_served=req.industries_served,
            revenue_range=req.revenue_range,
            employee_count=req.employee_count,
            key_stakeholders=req.key_stakeholders,
            change_readiness=req.change_readiness,
        )),
        "findings": (generate_findings, dict(brief, pillars=pillars_raw)),
        "risks": (generate_risks, dict(brief, pillars=pillars_raw)),
        "deep_dives": (generate_deep_dives, dict(brief, pillars=pillars_raw)),
        "architecture": (generate_architecture, brief),
        "plan": (generate_plan, dict(
            brief, pillars=pillars_raw, priority_areas=req.priority_areas,
        )),
        "granuler": (generate_granuler, dict(
            brief, granuler_location=req.granuler_location, locations=req.locations,
        )),
    }
    # A new client has no track record, so the slide that reports one is not
    # generated at all unless the assessor supplied the work.
    if req.prior_work.strip():
        jobs["prior_work"] = (generate_prior_work, dict(
            brief, prior_work=req.prior_work,
            savings_identified=req.savings_identified,
        ))

    with ThreadPoolExecutor(max_workers=20) as pool:
        futures = {name: pool.submit(fn, **kwargs) for name, (fn, kwargs) in jobs.items()}
        pillar_futures = [
            pool.submit(
                generate_pillar_content,
                company_name=_LLM_CLIENT_LABEL,
                pillar_name=p["pillar"],
                pillar_score=pillar_summaries[i]["score"],
                subtopics=p["subtopics"],
            )
            for i, p in enumerate(pillars_raw)
        ]
        content = {name: _restore_name(f.result(), req.company_name)
                   for name, f in futures.items()}
        content["pillars"] = [_restore_name(f.result(), req.company_name)
                              for f in pillar_futures]

    pptx_bytes, _slots = build_deck(intake, pillars_raw, content, SUBTOPICS_PER_PILLAR)

    filename = f"{req.company_name.replace(' ', '_')}_Granuler_Assessment.pptx"
    return Response(
        content=pptx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _deliverable_response(kind: str, company_name: str, result: dict, fmt: str):
    """JSON by default so the existing HTML output panels keep working."""
    if fmt != "pdf":
        return {"company_name": company_name, **result}

    render, label = PDF_RENDERERS[kind]
    filename = f"{company_name.replace(' ', '_')}_{label}.pdf"
    return Response(
        content=render(company_name, result),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


class NotesRequest(BaseModel):
    notes: str
    company_name: str


def _mask_company(text: str, company_name: str) -> str:
    """Replace the client's name in freeform notes with the LLM placeholder.

    The settled rule is that the client's company name never reaches the LLM.
    Notes extraction has to send the notes body, so the name is masked here
    first: the full name, then each distinctive word of it, longest first so a
    multi-word name is not left half-substituted.

    Person names in the notes are NOT masked - there is no reliable detector
    for them. The frontend states this above the notes box so the choice is the
    assessor's, made knowingly.
    """
    tokens = [company_name] + [
        word for word in re.split(r"\W+", company_name) if len(word) > 3
    ]
    for token in sorted(set(tokens), key=len, reverse=True):
        text = re.sub(
            rf"(?<![\w]){re.escape(token)}(?![\w])", _LLM_CLIENT_LABEL, text, flags=re.IGNORECASE
        )
    return text


@app.post("/extract-from-notes")
def extract_notes(req: NotesRequest):
    if not req.notes.strip():
        raise HTTPException(status_code=422, detail="notes must not be empty")
    if not req.company_name.strip():
        raise HTTPException(
            status_code=422,
            detail="company_name is required so it can be masked before the notes are sent",
        )

    masked = _mask_company(req.notes, req.company_name)
    result = _restore_name(
        extract_from_notes(company_name=_LLM_CLIENT_LABEL, notes=masked),
        req.company_name,
    )
    return {"company_name": req.company_name, **result}


class RenderPdfRequest(BaseModel):
    kind: str
    company_name: str
    data: dict


@app.post("/render-pdf")
def render_pdf(req: RenderPdfRequest):
    """Render an already-generated deliverable to PDF. No LLM call.

    The frontend holds the JSON it just displayed, so downloading a PDF of it
    must not re-run generation: that would cost another 20-40s and another set
    of tokens, and could return different text from what is on screen.
    """
    if req.kind not in PDF_RENDERERS:
        raise HTTPException(status_code=422, detail=f"Unknown deliverable {req.kind!r}")
    render, label = PDF_RENDERERS[req.kind]
    filename = f"{req.company_name.replace(' ', '_')}_{label}.pdf"
    return Response(
        content=render(req.company_name, req.data),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _parse_request(req: ReportRequest):
    pillars_raw = [p.model_dump() for p in req.pillars]
    intake = req.model_dump(exclude={"pillars"})
    pillar_summaries = [
        {"name": p["pillar"], "score": pillar_score(p["subtopics"], SUBTOPICS_PER_PILLAR)}
        for p in pillars_raw
    ]
    score = overall_score(pillars_raw, SUBTOPICS_PER_PILLAR)
    return pillars_raw, intake, pillar_summaries, score, band_for(score)[0]


@app.post("/generate-quick-wins")
def quick_wins(req: ReportRequest, format: str = Query("json", pattern="^(json|pdf)$")):
    if len(req.pillars) != PILLAR_COUNT:
        raise HTTPException(status_code=422, detail=f"Exactly {PILLAR_COUNT} pillars required")
    pillars_raw, _, _, overall_score, maturity_band = _parse_request(req)
    result = _restore_name(generate_quick_wins(
        company_name=_LLM_CLIENT_LABEL,
        industry=req.industry,
        business_goals=req.business_goals,
        pain_points=req.pain_points,
        pillars=pillars_raw,
        overall_score=overall_score,
        maturity_band=maturity_band,
    ), req.company_name)
    return _deliverable_response("quick-wins", req.company_name, result, format)


@app.post("/generate-risk-register")
def risk_register(req: ReportRequest, format: str = Query("json", pattern="^(json|pdf)$")):
    if len(req.pillars) != PILLAR_COUNT:
        raise HTTPException(status_code=422, detail=f"Exactly {PILLAR_COUNT} pillars required")
    pillars_raw, _, _, overall_score, maturity_band = _parse_request(req)
    result = _restore_name(generate_risk_register(
        company_name=_LLM_CLIENT_LABEL,
        industry=req.industry,
        pillars=pillars_raw,
        overall_score=overall_score,
        maturity_band=maturity_band,
    ), req.company_name)
    return _deliverable_response("risk-register", req.company_name, result, format)


@app.post("/generate-proposal")
def proposal(req: ReportRequest, format: str = Query("json", pattern="^(json|pdf)$")):
    if len(req.pillars) != PILLAR_COUNT:
        raise HTTPException(status_code=422, detail=f"Exactly {PILLAR_COUNT} pillars required")
    pillars_raw, _, pillar_summaries, overall_score, maturity_band = _parse_request(req)
    result = _restore_name(generate_proposal(
        company_name=_LLM_CLIENT_LABEL,
        industry=req.industry,
        overall_score=overall_score,
        maturity_band=maturity_band,
        business_goals=req.business_goals,
        pain_points=req.pain_points,
        major_risks=req.major_risks,
        founder_dependency=req.founder_dependency,
        budget_appetite=req.budget_appetite,
        pillar_summaries=pillar_summaries,
    ), req.company_name)
    return _deliverable_response("proposal", req.company_name, result, format)
