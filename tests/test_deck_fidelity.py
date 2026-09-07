"""Regression lock: no client's deck may carry another client's facts.

On 2026-08-31 a generated deck shipped 33 unedited template slides naming a
different client. The template is gone now - every slide is drawn from the
request - but the failure mode it protected against is not, so the guards move
here: nothing reaches a slide that did not come from this request, the LLM
placeholder never survives, and a missing content block removes its slide
rather than leaving a half-filled one.
"""
import io
import sys
from pathlib import Path

import pytest
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).parent.parent))

from api.deck import band_for, overall_score  # noqa: E402
from api.image_brief import slots_from_deck  # noqa: E402
from api.slides import build_deck  # noqa: E402
from tests import fixtures as fx  # noqa: E402

# Facts belonging to the other clients whose decks exist in this repo's history.
FOREIGN = ("Uni-tech", "Unitech", "Dhruv", "SAP", "S/4HANA", "Sydler", "Pune",
           "NIVDAS 2.1.1 ->5.0", "XYZ")


def _deck(content=None, pillars=None):
    data, _slots = build_deck(fx.INTAKE, pillars or fx.PILLARS_RAW,
                              content if content is not None else fx.CONTENT)
    return Presentation(io.BytesIO(data))


def _all_text(prs) -> str:
    out = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    out.extend(cell.text for cell in row.cells)
    return "\n".join(out)


@pytest.fixture(scope="module")
def deck():
    return _deck()


def test_deck_length_is_in_the_presentable_range(deck):
    assert 30 <= len(deck.slides._sldIdLst) <= 80


def test_no_foreign_client_appears_anywhere(deck):
    text = _all_text(deck)
    for term in FOREIGN:
        assert term not in text, f"another client's fact reached the deck: {term}"


def test_the_llm_placeholder_never_survives(deck):
    assert "the client company" not in _all_text(deck).lower()


def test_the_company_and_its_score_are_on_the_deck(deck):
    text = _all_text(deck)
    score = overall_score(fx.PILLARS_RAW, 4)
    assert fx.INTAKE["company_name"] in text
    assert f"{score:g}" in text
    assert band_for(score)[0] in text


def test_every_scored_subtopic_reaches_the_heatmap(deck):
    text = _all_text(deck)
    for pillar in fx.PILLARS_RAW:
        for sub in pillar["subtopics"]:
            assert sub["subtopic"] in text


def test_missing_content_removes_slides_rather_than_inventing_them():
    """An empty content block must cost slides, never fill them with defaults."""
    full = len(_deck().slides._sldIdLst)
    stripped = {k: ({} if k != "pillars" else v) for k, v in fx.CONTENT.items()}
    thin = len(_deck(stripped).slides._sldIdLst)
    assert thin < full
    assert "the client company" not in _all_text(_deck(stripped)).lower()


def test_an_inapplicable_deep_dive_is_dropped():
    content = {**fx.CONTENT, "deep_dives": {
        k: {**v, "applicable": False} for k, v in fx.CONTENT["deep_dives"].items()
    }}
    assert len(_deck(content).slides._sldIdLst) < len(_deck().slides._sldIdLst)


def test_every_image_placeholder_carries_its_prompt():
    data, slots = build_deck(fx.INTAKE, fx.PILLARS_RAW, fx.CONTENT)
    assert slots, "the deck opened no image slots"
    recovered = slots_from_deck(Presentation(io.BytesIO(data)))
    assert [s["n"] for s in recovered] == [s["n"] for s in slots]
    for original, found in zip(slots, recovered):
        assert found["prompt"] == original["prompt"].strip()
        assert found["px"] == original["px"]


def test_scores_change_the_deck():
    """An untouched form and a real assessment must not produce the same deck."""
    flat = [
        {"pillar": p["pillar"],
         "subtopics": [{**s, "score": 3} for s in p["subtopics"]]}
        for p in fx.PILLARS_RAW
    ]
    assert _all_text(_deck(pillars=flat)) != _all_text(_deck())
